from pathlib import Path
from textwrap import wrap
from html import escape
import subprocess

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.document import Document as _Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image as RLImage, PageBreak, PageTemplate,
    Paragraph as RLParagraph, Spacer, Table as RLTable, TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "DeepRevision_Agent架构与协作说明.docx"
ASSET_DIR = ROOT / "artifacts" / "agent_architecture_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

FONT_REGULAR = "/System/Library/Fonts/STHeiti Light.ttc"
FONT_BOLD = "/System/Library/Fonts/STHeiti Medium.ttc"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F4F7FB"
GRAY = "666666"
LIGHT_GRAY = "D9E0E8"
WHITE = "FFFFFF"
BLACK = "000000"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def text_bbox(draw, text, f):
    box = draw.textbbox((0, 0), text, font=f)
    return box[2] - box[0], box[3] - box[1]


def center_text(draw, box, lines, size=32, bold=False, fill="#17365D", spacing=8):
    f = font(size, bold)
    if isinstance(lines, str):
        lines = lines.split("\n")
    heights = [text_bbox(draw, line, f)[1] for line in lines]
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total_h) / 2
    for line, h in zip(lines, heights):
        w, _ = text_bbox(draw, line, f)
        x = box[0] + (box[2] - box[0] - w) / 2
        draw.text((x, y), line, font=f, fill=fill)
        y += h + spacing


def rounded_box(draw, box, title, subtitle=None, fill="#F4F7FB", outline="#A9B7C7", width=3,
                title_size=34, subtitle_size=25):
    draw.rounded_rectangle(box, radius=20, fill=fill, outline=outline, width=width)
    if subtitle:
        mid = box[1] + (box[3] - box[1]) * 0.42
        center_text(draw, (box[0] + 10, box[1] + 8, box[2] - 10, mid + 12), title,
                    title_size, True, "#17365D")
        center_text(draw, (box[0] + 18, mid, box[2] - 18, box[3] - 8), subtitle,
                    subtitle_size, False, "#444444", 6)
    else:
        center_text(draw, box, title, title_size, True, "#17365D")


def arrow(draw, start, end, color="#59789C", width=5, label=None, label_offset=(0, -28)):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    for delta in (2.55, -2.55):
        p = (end[0] + length * math.cos(angle + delta),
             end[1] + length * math.sin(angle + delta))
        draw.line([end, p], fill=color, width=width)
    if label:
        f = font(23, False)
        w, h = text_bbox(draw, label, f)
        x = (start[0] + end[0]) / 2 - w / 2 + label_offset[0]
        y = (start[1] + end[1]) / 2 - h / 2 + label_offset[1]
        draw.rounded_rectangle((x - 8, y - 4, x + w + 8, y + h + 4), 8, fill="#FFFFFF")
        draw.text((x, y), label, font=f, fill="#444444")


def make_architecture_diagram(path):
    img = Image.new("RGB", (1900, 1320), "white")
    d = ImageDraw.Draw(img)
    center_text(d, (0, 20, 1900, 90), "DeepRevision 分层多 Agent 架构", 46, True, "#000000")

    rounded_box(d, (710, 110, 1190, 230), "用户与 Next.js 前端", "输入问题、上传课件、接收 SSE 事件", "#EAF2F8")
    rounded_box(d, (710, 290, 1190, 410), "FastAPI 接口层", "鉴权、会话归属、幂等、参数校验", "#EAF2F8")
    rounded_box(d, (710, 470, 1190, 610), "Supervisor", "意图识别、参数提取、选择主 Agent", "#DCE6F1")
    rounded_box(d, (710, 670, 1190, 810), "ToolOrchestrator", "工具规划、权限分级、确认、失败控制", "#DCE6F1")

    arrow(d, (950, 230), (950, 290))
    arrow(d, (950, 410), (950, 470), label="加载 checkpoint")
    arrow(d, (950, 610), (950, 670), label="route + route_params")

    agents = [
        ("RAGAgent", "课件问答与证据引用"),
        ("QuizAgent", "个性化小测与批改"),
        ("ExamAgent", "分阶段生成完整试卷"),
        ("PlannerAgent", "制定约束化复习计划"),
        ("HistoryAgent", "查询错题与掌握度"),
        ("OpsAgent", "文件、会话与导出操作"),
        ("LearningLoop", "跨请求学习闭环"),
        ("Chitchat", "普通对话"),
    ]
    xs = [55, 505, 955, 1405]
    for i, (name, desc) in enumerate(agents):
        row = i // 4
        col = i % 4
        x = xs[col]
        y = 900 + row * 170
        rounded_box(d, (x, y, x + 390, y + 120), name, desc, "#F4F7FB",
                    title_size=31, subtitle_size=24)
        start_x = 810 + col * 90
        arrow(d, (950, 810), (x + 195, y), width=4)

    center_text(d, (20, 1240, 1880, 1300),
                "一次请求通常只选择一个主业务 Agent；跨 Agent 协作主要由 LearningLoop 组织。",
                28, False, "#444444")
    img.save(path, dpi=(220, 220))


def make_collaboration_diagram(path):
    img = Image.new("RGB", (1900, 1000), "white")
    d = ImageDraw.Draw(img)
    center_text(d, (0, 20, 1900, 90), "Agent 之间如何传递信息", 46, True, "#000000")

    rounded_box(d, (90, 150, 430, 285), "用户输入", "自然语言和客户端动作", "#EAF2F8")
    rounded_box(d, (580, 130, 1320, 305), "SupervisorState", ["input、session_id、route、route_params", "memory_context、tool_context、subagent_result"], "#DCE6F1")
    rounded_box(d, (1470, 150, 1810, 285), "业务 Agent", "读取状态并写回结构化结果", "#EAF2F8")
    arrow(d, (430, 218), (580, 218))
    arrow(d, (1320, 218), (1470, 218))

    services = [
        ((120, 460, 500, 610), "Memory 服务", "消息、错题、掌握度"),
        ((550, 460, 930, 610), "RAG 服务", "混合检索、重排、证据"),
        ((980, 460, 1360, 610), "工具服务", "查询、导出、管理操作"),
        ((1410, 460, 1790, 610), "模型服务", "生成、分类、质量检查"),
    ]
    for box, title, subtitle in services:
        rounded_box(d, box, title, subtitle, "#F4F7FB", title_size=31, subtitle_size=24)
        arrow(d, ((box[0] + box[2]) // 2, box[1]), (950, 305), width=4)

    rounded_box(d, (140, 760, 600, 900), "LangGraph Checkpoint", "保存图执行位置和请求状态", "#F5F5F5")
    rounded_box(d, (720, 760, 1180, 900), "SQLite 业务状态", "学习阶段、练习记录、掌握度", "#F5F5F5")
    rounded_box(d, (1300, 760, 1760, 900), "ChromaDB", "向量、父子块和文档元数据", "#F5F5F5")
    arrow(d, (820, 305), (370, 760), label="流程恢复", label_offset=(-40, 0))
    arrow(d, (950, 305), (950, 760), label="跨请求记忆", label_offset=(55, 0))
    arrow(d, (1080, 305), (1530, 760), label="检索依据", label_offset=(40, 0))
    img.save(path, dpi=(220, 220))


def make_learning_loop_diagram(path):
    img = Image.new("RGB", (1900, 1050), "white")
    d = ImageDraw.Draw(img)
    center_text(d, (0, 20, 1900, 90), "LearningLoop 持续学习闭环", 46, True, "#000000")

    nodes = [
        ("学情诊断", "History 与 Memory\n读取错题和掌握度", 160, 180),
        ("课件取证", "RAG 检索薄弱点\n对应的课件内容", 720, 180),
        ("制定计划", "Planner 生成多日任务\n显式约束强制覆盖", 1280, 180),
        ("当日出题", "调用 QuizAgent\nGenerate Critic Revise", 1280, 620),
        ("作答评估", "保存练习记录\n更新知识点掌握度", 720, 620),
        ("推进状态", "下一日 补弱 或完成\n写入 agent_states", 160, 620),
    ]
    boxes = []
    for title, subtitle, x, y in nodes:
        box = (x, y, x + 440, y + 170)
        boxes.append(box)
        rounded_box(d, box, title, subtitle, "#F4F7FB", title_size=35, subtitle_size=26)
    arrow(d, (600, 265), (720, 265), label="薄弱点")
    arrow(d, (1160, 265), (1280, 265), label="证据")
    arrow(d, (1500, 350), (1500, 620), label="计划参数", label_offset=(70, 0))
    arrow(d, (1280, 705), (1160, 705), label="题目")
    arrow(d, (720, 705), (600, 705), label="结果")
    arrow(d, (380, 620), (380, 350), label="继续周期", label_offset=(-75, 0))
    center_text(d, (200, 900, 1700, 990),
                "LearningLoop 负责推进长期任务；QuizAgent 负责一次出题，二者职责不同。",
                30, False, "#444444")
    img.save(path, dpi=(220, 220))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="STHeiti", east_asia="STHeiti", size=11, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, bold_lead=None, space_after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    for run in p.runs:
        set_run_font(run)
    if not p.runs:
        set_run_font(p.add_run(text))
    else:
        p.runs[0].text = text
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0])
    else:
        set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        set_run_font(r, east_asia="Microsoft YaHei", size=16 if level == 1 else 13,
                     bold=True, color=BLACK)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].width = Inches(widths[i])
        set_cell_shading(hdr[i], NAVY)
        set_cell_border(hdr[i])
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(h), size=10, bold=True, color=WHITE)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            set_cell_shading(cells[i], WHITE if idx % 2 == 0 else PALE)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(str(value)), size=9.5, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_picture(doc, path, width=6.8, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(7)
        set_run_font(c.add_run(caption), size=9, color=GRAY)


def iter_blocks(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, parent)


def build_stable_pdf(source_doc, pdf_path):
    """Reflow the DOCX content with an embedded CJK font for deterministic QA."""
    pdfmetrics.registerFont(TTFont("STHeitiEmbed", FONT_REGULAR))
    pdfmetrics.registerFont(TTFont("STHeitiEmbedBold", FONT_BOLD))

    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "BodyCN", parent=base["BodyText"], fontName="STHeitiEmbed", fontSize=10.4,
        leading=16, alignment=TA_JUSTIFY, textColor=colors.black,
        spaceAfter=6, wordWrap="CJK",
    )
    title = ParagraphStyle(
        "TitleCN", parent=body, fontName="STHeitiEmbedBold", fontSize=24,
        leading=32, alignment=TA_CENTER, spaceBefore=150, spaceAfter=24,
    )
    h1 = ParagraphStyle(
        "H1CN", parent=body, fontName="STHeitiEmbedBold", fontSize=16,
        leading=22, spaceBefore=13, spaceAfter=8, keepWithNext=True,
    )
    h2 = ParagraphStyle(
        "H2CN", parent=body, fontName="STHeitiEmbedBold", fontSize=13,
        leading=18, spaceBefore=9, spaceAfter=6, keepWithNext=True,
    )
    center = ParagraphStyle("CenterCN", parent=body, alignment=TA_CENTER, spaceAfter=10)
    caption = ParagraphStyle(
        "CaptionCN", parent=body, alignment=TA_CENTER, fontSize=8.5,
        leading=12, textColor=colors.HexColor("#666666"), spaceAfter=8,
    )
    bullet = ParagraphStyle(
        "BulletCN", parent=body, leftIndent=18, firstLineIndent=-10, spaceAfter=3,
    )

    frame = Frame(0.72 * inch, 0.68 * inch, 7.06 * inch, 9.62 * inch, id="normal")

    def draw_footer(canvas, doc_template):
        canvas.saveState()
        canvas.setFont("STHeitiEmbed", 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawCentredString(4.25 * inch, 0.38 * inch,
                                 f"DeepRevision Agent 架构与协作说明  {doc_template.page}")
        canvas.restoreState()

    pdf_doc = BaseDocTemplate(str(pdf_path), pagesize=letter,
                              leftMargin=0.72 * inch, rightMargin=0.72 * inch,
                              topMargin=0.68 * inch, bottomMargin=0.68 * inch,
                              title="DeepRevision Agent 架构与协作说明")
    pdf_doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=draw_footer)])

    story = []
    image_index = 0
    number_counter = 0
    for block in iter_blocks(source_doc):
        if isinstance(block, DocxParagraph):
            brs = block._p.findall(".//" + qn("w:br"))
            if any(br.get(qn("w:type")) == "page" for br in brs):
                story.append(PageBreak())
                number_counter = 0
                continue

            blips = block._p.findall(".//" + qn("a:blip"))
            if blips:
                rid = blips[0].get(qn("r:embed"))
                part = source_doc.part.related_parts[rid]
                ext = Path(part.partname).suffix or ".png"
                extracted = ASSET_DIR / f"embedded_{image_index}{ext}"
                extracted.write_bytes(part.blob)
                image_index += 1
                with Image.open(extracted) as im:
                    ratio = im.height / im.width
                target_w = 6.75 * inch
                target_h = target_w * ratio
                if target_h > 7.5 * inch:
                    target_h = 7.5 * inch
                    target_w = target_h / ratio
                story.append(RLImage(str(extracted), width=target_w, height=target_h, hAlign="CENTER"))
                story.append(Spacer(1, 4))
                continue

            text = block.text.strip()
            if not text:
                story.append(Spacer(1, 5))
                continue
            style_name = block.style.name if block.style else "Normal"
            safe = escape(text).replace("\n", "<br/>")
            if style_name == "Title":
                story.append(RLParagraph(safe, title))
            elif style_name == "Heading 1":
                number_counter = 0
                story.append(RLParagraph(safe, h1))
            elif style_name in ("Heading 2", "Heading 3"):
                number_counter = 0
                story.append(RLParagraph(safe, h2))
            elif style_name.startswith("List Bullet"):
                story.append(RLParagraph("• " + safe, bullet))
            elif style_name.startswith("List Number"):
                number_counter += 1
                story.append(RLParagraph(f"{number_counter}. {safe}", bullet))
            elif text.startswith("图 "):
                story.append(RLParagraph(safe, caption))
            elif style_name == "Subtitle":
                story.append(RLParagraph(safe, center))
            else:
                number_counter = 0
                story.append(RLParagraph(safe, body))
        else:
            data = []
            for r_idx, row in enumerate(block.rows):
                row_data = []
                for cell in row.cells:
                    content = escape("\n".join(p.text for p in cell.paragraphs)).replace("\n", "<br/>")
                    cell_style = ParagraphStyle(
                        f"cell-{len(data)}-{len(row_data)}", parent=body,
                        fontName="STHeitiEmbedBold" if r_idx == 0 else "STHeitiEmbed",
                        fontSize=8.7, leading=12, alignment=TA_CENTER if r_idx == 0 else TA_LEFT,
                        textColor=colors.white if r_idx == 0 else colors.black,
                    )
                    row_data.append(RLParagraph(content, cell_style))
                data.append(row_data)
            cols = len(data[0]) if data else 1
            widths = [6.8 * inch / cols] * cols
            if cols == 2:
                widths = [2.2 * inch, 4.6 * inch]
            elif cols == 3:
                widths = [1.35 * inch, 2.55 * inch, 2.9 * inch]
            tbl = RLTable(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
            style_cmds = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
            for row_idx in range(1, len(data)):
                if row_idx % 2 == 0:
                    style_cmds.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#F4F7FB")))
            tbl.setStyle(TableStyle(style_cmds))
            story.append(tbl)
            story.append(Spacer(1, 8))
            number_counter = 0

    pdf_doc.build(story)


def wrap_pdf_pages_as_docx(pdf_path, output_path):
    page_prefix = ASSET_DIR / "final_page"
    for old in ASSET_DIR.glob("final_page-*.png"):
        old.unlink()
    subprocess.run([
        "pdftoppm", "-png", "-r", "180", str(pdf_path), str(page_prefix)
    ], check=True)
    pages = sorted(ASSET_DIR.glob("final_page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    final_doc = Document()
    sec = final_doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.1)
    sec.bottom_margin = Inches(0.1)
    sec.left_margin = Inches(0.1)
    sec.right_margin = Inches(0.1)
    normal = final_doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    for idx, page in enumerate(pages):
        p = final_doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(page), width=Inches(8.3))
        if idx < len(pages) - 1:
            p.add_run().add_break()
            p.runs[-1]._element.findall(".//" + qn("w:br"))[0].set(qn("w:type"), "page")
    final_doc.core_properties.title = "DeepRevision Agent 架构与协作说明"
    final_doc.core_properties.subject = "AI 应用与 Agent 后端岗位面试准备"
    final_doc.core_properties.author = "熊浩宇"
    final_doc.save(output_path)


def build_document():
    arch = ASSET_DIR / "architecture.png"
    collab = ASSET_DIR / "collaboration.png"
    loop = ASSET_DIR / "learning_loop.png"
    make_architecture_diagram(arch)
    make_collaboration_diagram(collab)
    make_learning_loop_diagram(loop)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "STHeiti"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "STHeiti")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "STHeiti")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        st = styles[style_name]
        st.font.name = "STHeiti"
        st._element.rPr.rFonts.set(qn("w:ascii"), "STHeiti")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "STHeiti")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
        st.font.color.rgb = RGBColor(0, 0, 0)

    # Remove the built-in Title border that LibreOffice may render as a blue rule.
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    # Cover page
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DeepRevision Agent 架构与协作说明")
    set_run_font(r, size=26, bold=True)
    p.style = styles["Title"]

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(18)
    p2.paragraph_format.space_after = Pt(30)
    set_run_font(p2.add_run("面试理解与讲解手册"), size=15, color=GRAY)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.left_indent = Inches(0.55)
    intro.paragraph_format.right_indent = Inches(0.55)
    intro.paragraph_format.line_spacing = 1.35
    intro.paragraph_format.space_before = Pt(24)
    set_run_font(intro.add_run(
        "本文档说明 DeepRevision 中各类 Agent 的职责、协作方式、状态传递和持久化机制。"
        "系统采用 Supervisor 主导的分层多 Agent 架构。Supervisor 决定由谁处理请求，"
        "ToolOrchestrator 控制工具调用，业务 Agent 完成具体任务，RAG、Memory 和数据库提供共享能力。"
    ), size=12)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(70)
    set_run_font(p3.add_run("适用场景  AI 应用与 Agent 后端岗位面试"), size=11, color=GRAY)
    doc.add_page_break()

    add_heading(doc, "一 系统定位", 1)
    add_para(doc,
             "DeepRevision 是一个 Supervisor 统一编排的分层多 Agent 系统。业务 Agent 不进行无约束的自由对话。"
             "一次普通请求通常只选择一个主 Agent，复杂的长期学习任务由 LearningLoop 组织多个能力共同完成。")
    add_para(doc, "面试时应先讲清三层分工：")
    add_bullet(doc, "Supervisor 判断请求属于什么任务，并提取 topic、题量、题型等参数。")
    add_bullet(doc, "ToolOrchestrator 判断需要调用哪些工具，并执行权限、确认、重试和调用次数控制。")
    add_bullet(doc, "业务 Agent 组合 RAG、Memory、模型和数据库，产出结构化结果。")
    add_picture(doc, arch, 6.8, "图 1  DeepRevision 分层多 Agent 架构")

    add_heading(doc, "二 顶层 Agent 职责", 1)
    role_rows = [
        ("Supervisor", "识别意图、抽取参数、选择主 Agent，并解释路由原因。", "位于所有业务 Agent 之前；写入 route 和 route_params。"),
        ("ToolOrchestrator", "规划工具调用，执行权限分级、危险操作确认、失败处理和最大步数限制。", "把工具 observation 写入 tool_context，供业务 Agent 复用。"),
        ("RAGAgent", "回答课件知识问题，组织检索上下文、生成答案并返回可验证引用。", "调用 RAG 服务和模型服务；检索失败时不伪造证据。"),
        ("QuizAgent", "根据课件、错题和学习画像生成小测，也处理答题后的批改。", "读取 Memory 和 RAG；内部运行 Generate Critic Revise。"),
        ("ExamAgent", "按题型和数量分阶段生成完整试卷，合并、去重并校验总分。", "调用 RAG；每个阶段运行生成、审查和修订子图。"),
        ("PlannerAgent", "根据时间约束、学习目标和掌握度生成复习计划。", "读取学习画像；用户明确的天数和时长覆盖模型输出。"),
        ("HistoryAgent", "查询练习历史、错题、正确率和知识点掌握度。", "直接读取 SQLite 中的事实数据，避免由模型编造统计。"),
        ("OpsAgent", "处理文件、会话、课件、导出、诊断和联网查询等操作。", "依赖工具层；删除等危险写操作必须获得明确确认。"),
        ("LearningLoop", "组织诊断、计划、出题、评估和补弱，推进跨请求学习周期。", "复用 History、RAG 和 Planner 能力，并显式调用 QuizAgent。"),
        ("Chitchat", "处理问候、闲聊和不需要业务能力的普通对话。", "使用必要的消息历史，不触发无意义的检索或工具调用。"),
    ]
    add_table(doc, ["组件", "主要职责", "与其他组件的关系"], role_rows, [1.25, 2.72, 2.97])

    add_heading(doc, "三 内部节点和共享服务", 1)
    add_para(doc,
             "并非所有带名称的模块都是顶层 Agent。Generate、Critic 和 Revise 是 QuizAgent 或 ExamAgent 内部的工作流节点；"
             "RAG、Memory、数据库和工具函数是共享服务。面试时区分这些概念，可以避免把普通函数包装成多 Agent。")
    internal_rows = [
        ("Generate", "按照结构化约束生成题目、试卷或计划候选结果。"),
        ("Critic", "检查考点覆盖、答案一致性、证据支持、重复题和格式完整性。"),
        ("Revise", "根据 Critic 的问题清单定向修订；达到次数上限后停止。"),
        ("RAG 服务", "执行查询改写、BM25 与向量召回、RRF 融合、重排、父块展开和证据提取。"),
        ("Memory 服务", "提供近期消息、长期知识、错题、练习记录和知识点掌握度。"),
        ("工具集合", "封装检索、统计、会话管理、导出和外部查询等可执行操作。"),
    ]
    add_table(doc, ["模块", "作用"], internal_rows, [1.55, 5.35])

    add_heading(doc, "四 Agent 之间如何通信", 1)
    add_para(doc,
             "Agent 之间主要通过统一的 SupervisorState 传递结构化数据，而不是互相发送自由文本。"
             "这使路由、参数和工具结果可以被校验、记录和恢复。")
    add_picture(doc, collab, 6.8, "图 2  状态和共享服务之间的数据流")

    state_rows = [
        ("input", "当前用户输入。"),
        ("session_id", "当前会话标识，用于会话归属和数据隔离。"),
        ("route", "Supervisor 选择的主业务 Agent。"),
        ("route_params", "topic、count、question_types 等结构化业务参数。"),
        ("memory_context", "近期消息、长期记忆和学习画像。"),
        ("tool_context", "工具调用产生的 observation，可供后续节点复用。"),
        ("subagent_result", "业务 Agent 返回的结构化结果。"),
        ("final_answer", "序列化为 SSE 事件并返回前端的最终内容。"),
    ]
    add_table(doc, ["字段", "含义"], state_rows, [1.65, 5.25])

    add_heading(doc, "五 QuizAgent 的内部流程", 1)
    add_number(doc, "读取用户明确指定的题量、考点和题型。")
    add_number(doc, "读取最近错题、知识点掌握度和近期题目签名。")
    add_number(doc, "从课件中检索与目标知识点相关的原始证据。")
    add_number(doc, "Generate 节点根据约束和证据生成结构化题目。")
    add_number(doc, "Critic 检查考点覆盖、答案唯一性、解析一致性、证据支持和重复情况。")
    add_number(doc, "未通过时由 Revise 定向修改，再次进入 Critic；修订次数受上限限制。")
    add_number(doc, "Pydantic 执行字段和类型校验，本地规则完成去重和答案一致性检查。")
    add_number(doc, "QuizSet 通过 SSE 事件返回前端，渲染为交互式答题卡。")
    add_para(doc, "Critic 检查重点：", bold_lead="Critic 检查重点：")
    for item in [
        "题目是否覆盖用户要求的考点和难度。",
        "答案是否能够被课件证据支持。",
        "选择题是否存在多个合理答案。",
        "标准答案与解析是否一致。",
        "是否与最近生成的题目重复。",
        "Pydantic 结构是否满足前端渲染要求。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "六 ExamAgent 的内部流程", 1)
    add_para(doc,
             "ExamAgent 面向完整试卷。它先解析考点、题型、数量和分值，再按照选择题、填空判断题、主观题等阶段顺序执行。"
             "大阶段串行，阶段内部使用受限并发，防止同时生成过多题目导致重复、超时和模型限流。")
    exam_rows = [
        ("参数解析", "确定主题、题型分布、数量、难度和样卷格式。"),
        ("聚合检索", "从课件中构建动态考点池，并平衡不同主题的使用次数。"),
        ("阶段生成", "每个阶段运行 Generate Critic Revise，阶段失败允许有限重试。"),
        ("全卷合并", "统一编号、去重、检查题型数量并调整总分为 100。"),
        ("阶段补偿", "通过 exam_rerun_stage 和 exam_partial_questions 重跑失败阶段。"),
        ("结果返回", "输出 exam_paper 和 exam_canvas；导出由独立接口完成。"),
    ]
    add_table(doc, ["阶段", "处理内容"], exam_rows, [1.55, 5.35])

    add_heading(doc, "七 LearningLoop 如何组织多 Agent", 1)
    add_para(doc,
             "LearningLoop 是系统中最明显的跨能力协作流程。它保存当前学习阶段和下一步动作，因此用户不需要在每一轮重新说明学习目标。"
             "QuizAgent 负责一次出题，LearningLoop 负责决定何时出题、出什么题，以及作答后进入下一天还是补弱。")
    add_picture(doc, loop, 6.8, "图 3  LearningLoop 状态推进流程")

    phase_rows = [
        ("planned", "已完成诊断和计划，等待开始当日任务。"),
        ("dayN_quiz_issued", "第 N 天题目已发放，等待用户作答。"),
        ("dayN_answered", "已收到答案，等待评估和更新掌握度。"),
        ("dayN_reviewed", "当日评估完成，准备推进下一任务。"),
        ("remediation", "知识点未达到目标，进入针对性补弱。"),
        ("cycle_completed", "学习周期目标完成。"),
    ]
    add_table(doc, ["阶段", "含义"], phase_rows, [2.05, 4.85])

    add_heading(doc, "八 三类状态和恢复机制", 1)
    state_kind_rows = [
        ("SupervisorState", "当前输入、路由、工具结果和 Agent 输出。", "一次图执行过程。"),
        ("LangGraph Checkpoint", "图执行位置和对应 State。", "异常恢复、重复请求续跑和回档。"),
        ("agent_states", "LearningLoop 阶段、当前天数、计划和下一动作。", "跨多次用户请求长期保存。"),
        ("业务数据", "消息、练习记录、知识点掌握度和课件元数据。", "持续保存并供多个 Agent 查询。"),
    ]
    add_table(doc, ["状态类型", "保存内容", "解决的问题"], state_kind_rows, [1.65, 2.85, 2.4])
    add_para(doc,
             "Checkpoint 和 agent_states 的作用不同。Checkpoint 回答的是本次工作流执行到哪个节点；agent_states 回答的是用户的学习周期进行到哪一天。"
             "即使某次 HTTP 连接中断，业务学习阶段仍能通过 agent_states 恢复，但普通文本流不会从中断的 token 位置续写。")

    add_heading(doc, "九 示例请求的完整协作过程", 1)
    add_para(doc, "示例请求：根据我最近的错题，帮我针对薄弱知识点出 5 道题。", bold_lead="示例请求：")
    example_steps = [
        "前端把文本、session_id 和客户端动作提交到 FastAPI，并开始监听 SSE。",
        "接口层校验用户身份、会话归属和幂等键，然后加载对应 LangGraph checkpoint。",
        "Supervisor 将请求路由为 quiz，提取 count 等于 5，topic 标记为从学习画像动态解析。",
        "ToolOrchestrator 判断该请求只需要读操作，允许查询练习记录、掌握度和课件。",
        "QuizAgent 从 SQLite 读取最近错题和 knowledge_mastery，确定具体薄弱知识点。",
        "QuizAgent 调用 RAG 服务检索相关课件片段，经混合召回、融合和重排获得证据。",
        "内部 Generate Critic Revise 子图生成并校验 5 道题，Pydantic 验证结构。",
        "结果写入 subagent_result，接口依次发送路由状态、检索状态、题目卡片和完成事件。",
        "前端根据事件类型更新状态区、引用信息和答题组件。用户提交答案后，后端保存练习记录并更新掌握度。",
    ]
    for step in example_steps:
        add_number(doc, step)

    add_heading(doc, "十 工程限制和安全边界", 1)
    safety_rows = [
        ("路由限制", "route 使用白名单；结构化结果还要进行参数和语义漂移校验。"),
        ("工具权限", "工具分为只读、受控写、危险写、高成本导出和委托；危险写需要确认。"),
        ("调用上限", "ToolOrchestrator 和 Critic 修订都有最大步数，避免无限循环。"),
        ("结构化输出", "Pydantic 校验字段、类型和枚举，本地规则补充业务一致性检查。"),
        ("证据约束", "只有成功检索且能在上下文中匹配的原文才允许作为引用。"),
        ("幂等控制", "写操作结合用户、会话和请求键，防止重复提交造成重复数据。"),
        ("权限校验", "后端根据认证用户校验 session_id 归属，不信任前端传入的用户标识。"),
        ("失败降级", "模型调用设置超时、有限重试和备用模型；检索失败返回明确状态。"),
    ]
    add_table(doc, ["控制点", "工程手段"], safety_rows, [1.6, 5.3])

    add_heading(doc, "十一 面试回答", 1)
    add_para(doc,
             "DeepRevision 采用 Supervisor 主导的分层多 Agent 架构。Supervisor 通过规则预检和模型结构化分类识别意图，"
             "再由 ToolOrchestrator 完成工具规划、权限校验和失败控制，最后路由到 RAG、Quiz、Exam、Planner、History、Ops 或 LearningLoop。"
             "各 Agent 不进行无约束的自由对话，而是通过 LangGraph State、工具 observation 和持久化业务状态交换数据。"
             "Quiz 和 Exam 内部采用 Generate Critic Revise 质量闭环；LearningLoop 则串联学习画像、课件检索、复习计划、出题和结果评估，"
             "实现跨请求的持续学习。整个设计把模型用于需要语义判断和内容生成的环节，把权限、状态、幂等、约束覆盖和数据更新留给确定性代码。")

    add_heading(doc, "十二 容易混淆的问题", 1)
    qa_rows = [
        ("所有 Agent 会同时工作吗", "不会。普通请求通常只进入一个主业务 Agent；只有 LearningLoop 等复合流程会显式组合多个能力。"),
        ("RAG 是 Agent 吗", "RAG 服务本身是共享能力；RAGAgent 是使用该能力完成问答的业务 Agent。"),
        ("Critic 是顶层 Agent 吗", "不是。它是 Quiz、Exam 或 LearningLoop 内部的质量审查角色。"),
        ("State 能长期保存吗", "请求 State 由 checkpoint 保存；学习周期等业务状态由 agent_states 长期保存。"),
        ("子 Agent 会自由互调吗", "不会。调用关系由代码和图的条件边限制，结果通过结构化状态传递。"),
        ("为什么还需要 ToolOrchestrator", "Supervisor 解决任务归属；ToolOrchestrator 解决工具选择、权限、确认、重试和调用上限。"),
    ]
    add_table(doc, ["问题", "回答"], qa_rows, [2.0, 4.9])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("DeepRevision Agent 架构与协作说明"), size=9, color=GRAY)

    doc.core_properties.title = "DeepRevision Agent 架构与协作说明"
    doc.core_properties.subject = "AI 应用与 Agent 后端岗位面试准备"
    doc.core_properties.author = "熊浩宇"
    editable_source = ASSET_DIR / "editable_source.docx"
    doc.save(editable_source)
    stable_pdf = ASSET_DIR / "stable_layout.pdf"
    build_stable_pdf(doc, stable_pdf)
    wrap_pdf_pages_as_docx(stable_pdf, OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
