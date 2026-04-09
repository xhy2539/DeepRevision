"""
试卷导出 API
支持导出 Word 格式试卷和答案卷
"""
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel
from typing import Optional
import json
import re
import io
import uuid
import os

router = APIRouter(prefix="/api/exam", tags=["exam_export"])

# 临时文件目录
TEMP_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "data", "temp")
os.makedirs(TEMP_DIR, exist_ok=True)


class ExamFormatRequest(BaseModel):
    """试卷格式化请求"""
    exam_paper: str
    course_name: str = "期末考试"
    include_difficulty: bool = True


class ExamExportRequest(BaseModel):
    """试卷导出请求"""
    exam_paper: str
    course_name: str = "期末考试"
    include_answers: bool = True


class AnswerSheetRequest(BaseModel):
    """答案卷生成请求"""
    exam_paper: str
    course_name: str = "期末考试"


def parse_exam_content(exam_paper: str) -> dict:
    """
    解析试卷内容，提取题目、答案、解析
    返回结构化数据
    """
    result = {
        "title": "",
        "subtitle": "",
        "total_score": 100,
        "questions": []
    }

    lines = exam_paper.split('\n')
    current_section = None
    current_question = None

    # 题型映射 - 支持多种格式: "## 选择题", "1. 选择题", "【选择题】", "一、选择题"
    type_patterns = {
        '选择题': [re.compile(r'^##\s*选择题'), re.compile(r'^\d+[.、]\s*[\(（]?\s*选择题'), re.compile(r'【选择题】|一、选择题')],
        '填空题': [re.compile(r'^##\s*填空题'), re.compile(r'^\d+[.、]\s*[\(（]?\s*填空题'), re.compile(r'【填空题】|二、填空题')],
        '判断题': [re.compile(r'^##\s*判断题'), re.compile(r'^\d+[.、]\s*[\(（]?\s*判断题'), re.compile(r'【判断题】|三、判断题')],
        '简答题': [re.compile(r'^##\s*简答题'), re.compile(r'^\d+[.、]\s*[\(（]?\s*简答题'), re.compile(r'【简答题】|四、简答题')],
        '计算题': [re.compile(r'^##\s*计算题'), re.compile(r'^\d+[.、]\s*[\(（]?\s*计算题'), re.compile(r'【计算题】|五、计算题')],
        '分析题': [re.compile(r'^##\s*分析题'), re.compile(r'^\d+[.、]\s*[\(（]?\s*分析题'), re.compile(r'【分析题】|六、分析题')],
        '论述题': [re.compile(r'^##\s*论述题'), re.compile(r'^\d+[.、]\s*[\(（]?\s*论述题'), re.compile(r'【论述题】|七、论述题')],
        '名词解释': [re.compile(r'^##\s*名词解释'), re.compile(r'^\d+[.、]\s*[\(（]?\s*名词解释'), re.compile(r'【名词解释】')],
    }

    # 提取标题
    for line in lines[:5]:
        is_question_type = False
        for patterns in type_patterns.values():
            for p in patterns:
                if p.search(line):
                    is_question_type = True
                    break
            if is_question_type:
                break

        if line.strip() and not is_question_type:
            if not result["title"]:
                result["title"] = line.strip()
            elif "期末" in line or "试卷" in line:
                result["subtitle"] = line.strip()

    # 解析题目
    for line in lines:
        line = line.strip()
        if not line:
            continue

        # 检测题型（fix #11：用 matched 标志位，避免 current_section 非空后短路）
        matched = False
        for qtype, patterns in type_patterns.items():
            for pattern in patterns:
                if pattern.search(line):
                    current_section = qtype
                    current_question = None
                    matched = True
                    break
            if matched:
                break
        if matched:
            continue

        # 检测题目编号
        q_match = re.match(r'^(\d+)[.、]\s*(.+)', line)
        if q_match and current_section:
            q_num = q_match.group(1)
            q_content = q_match.group(2).strip()

            # 检查是否是选项
            if re.match(r'^[A-D][.、．]\s*', q_content):
                if current_question:
                    mm_opt = re.match(r'^([A-D])[.、．]\s*(.+)$', q_content)
                    if mm_opt and mm_opt.group(2).strip():
                        if "options" not in current_question:
                            current_question["options"] = []
                        current_question["options"].append(f"{mm_opt.group(1)}. {mm_opt.group(2).strip()}")
                continue

            # 新题目
            content_text = q_content
            options = []
            if current_section == "选择题" and re.search(r'[A-D][.、．]\s*', q_content):
                split_parts = [p.strip() for p in re.split(r'(?=[A-D][.、．]\s*)', q_content) if p.strip()]
                if split_parts:
                    first = split_parts[0]
                    if not re.match(r'^[A-D][.、．]\s*', first):
                        content_text = first
                        split_parts = split_parts[1:]
                    else:
                        content_text = ""
                    for part in split_parts:
                        mm = re.match(r'^([A-D])[.、．]\s*(.+)$', part)
                        if mm and mm.group(2).strip():
                            options.append(f"{mm.group(1)}. {mm.group(2).strip()}")

            current_question = {
                "number": int(q_num),
                "type": current_section,
                "content": content_text,
                "options": options,
                "answer": "",
                "analysis": "",
                "score": 10,  # 默认分值
                "difficulty": "中等"  # 默认难度
            }
            result["questions"].append(current_question)

        # 提取答案
        elif current_question and ("答案" in line or "答案：" in line):
            answer = re.sub(r'^答案[：:]\s*', '', line)
            current_question["answer"] = answer

        # 提取解析
        elif current_question and ("解析" in line or "解析：" in line):
            analysis = re.sub(r'^解析[：:]\s*', '', line)
            current_question["analysis"] = analysis
        elif current_question and current_section == "选择题" and re.match(r'^[A-D][.、．]\s*', line):
            mm = re.match(r'^([A-D])[.、．]\s*(.+)$', line)
            if mm and mm.group(2).strip():
                if "options" not in current_question:
                    current_question["options"] = []
                current_question["options"].append(f"{mm.group(1)}. {mm.group(2).strip()}")
        elif current_question and current_section == "选择题" and re.search(r'\b[A-D][.、．]\s+', line):
            matches = re.findall(r'([A-D])[.、．]\s*(.+?)(?=\s+[A-D][.、．]\s+|$)', line)
            if matches:
                if "options" not in current_question:
                    current_question["options"] = []
                for letter, text in matches:
                    text = (text or "").strip()
                    if text:
                        current_question["options"].append(f"{letter}. {text}")

    return result


@router.post("/format")
async def format_exam(request: ExamFormatRequest):
    """
    格式化试卷，添加难度、分值等元信息
    """
    try:
        parsed = parse_exam_content(request.exam_paper)

        # 添加难度评估（基于题目内容长度）
        for q in parsed["questions"]:
            content_len = len(q["content"])
            if content_len < 50:
                q["difficulty"] = "简单"
                q["score"] = 5
            elif content_len < 150:
                q["difficulty"] = "中等"
                q["score"] = 10
            else:
                q["difficulty"] = "较难"
                q["score"] = 15

        # 统计
        parsed["total_questions"] = len(parsed["questions"])
        parsed["question_types"] = {}

        for q in parsed["questions"]:
            qtype = q["type"]
            if qtype not in parsed["question_types"]:
                parsed["question_types"][qtype] = {"count": 0, "total_score": 0}
            parsed["question_types"][qtype]["count"] += 1
            parsed["question_types"][qtype]["total_score"] += q["score"]

        return {
            "success": True,
            "data": parsed
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/export/docx")
async def export_docx(request: ExamExportRequest):
    """
    导出 Word 格式试卷
    """
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # 标题
        title = doc.add_heading(request.course_name, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 副标题
        if request.include_answers:
            subtitle = doc.add_paragraph("（含答案）")
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # 解析试卷内容
        parsed = parse_exam_content(request.exam_paper)

        # 按题型分组输出
        current_type = None

        for q in parsed["questions"]:
            # 新题型标题
            if q["type"] != current_type:
                current_type = q["type"]
                doc.add_heading(current_type, level=2)

            # 题目内容
            p = doc.add_paragraph()
            p.add_run(f"{q['number']}. ").bold = True

            # 处理题目内容（包含选项等）
            content_lines = q["content"].split('\n')
            for i, line in enumerate(content_lines):
                if i == 0:
                    p.add_run(line)
                else:
                    p2 = doc.add_paragraph(line)
                    p2.paragraph_format.left_indent = Inches(0.5)

            for opt in q.get("options", []) or []:
                opt_p = doc.add_paragraph(str(opt))
                opt_p.paragraph_format.left_indent = Inches(0.5)

            # 答案（如果需要）
            if request.include_answers and q["answer"]:
                answer_p = doc.add_paragraph()
                answer_p.add_run(f"答案：{q['answer']}").font.color.rgb = RGBColor(0, 128, 0)

            # 解析
            if request.include_answers and q["analysis"]:
                analysis_p = doc.add_paragraph()
                analysis_p.add_run(f"解析：{q['analysis']}").font.color.rgb = RGBColor(128, 128, 128)

            doc.add_paragraph()

        # 保存文件
        file_id = str(uuid.uuid4())
        filename = f"{request.course_name}_{file_id[:8]}.docx"
        filepath = os.path.join(TEMP_DIR, filename)

        doc.save(filepath)

        return {
            "success": True,
            "download_url": f"/api/exam/download/{filename}",
            "filename": filename
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/answersheet")
async def generate_answer_sheet(request: AnswerSheetRequest):
    """
    生成答案卷
    """
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # 标题
        title = doc.add_heading(f"{request.course_name} - 答案卷", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # 解析试卷内容
        parsed = parse_exam_content(request.exam_paper)

        # 按题型分组
        questions_by_type = {}
        for q in parsed["questions"]:
            qtype = q["type"]
            if qtype not in questions_by_type:
                questions_by_type[qtype] = []
            questions_by_type[qtype].append(q)

        # 输出答案
        for qtype, questions in questions_by_type.items():
            doc.add_heading(qtype, level=2)

            for q in questions:
                p = doc.add_paragraph()
                p.add_run(f"{q['number']}. ").bold = True

                if q["answer"]:
                    p.add_run(q["answer"]).font.color.rgb = RGBColor(0, 128, 0)
                else:
                    p.add_run("（无）").font.color.rgb = RGBColor(128, 128, 128)

                # 添加分值
                p.add_run(f"  ({q['score']}分)")

        # 保存文件
        file_id = str(uuid.uuid4())
        filename = f"{request.course_name}_答案卷_{file_id[:8]}.docx"
        filepath = os.path.join(TEMP_DIR, filename)

        doc.save(filepath)

        return {
            "success": True,
            "download_url": f"/api/exam/download/{filename}",
            "filename": filename
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/download/{filename}")
async def download_file(filename: str):
    """下载文件（fix #1：路径遍历防护；fix #18：下载后自动删除临时文件）"""
    from fastapi.responses import FileResponse
    from starlette.background import BackgroundTask

    # 只取文件名部分，防止路径遍历
    safe_name = os.path.basename(filename)
    filepath = os.path.realpath(os.path.join(TEMP_DIR, safe_name))
    temp_dir_real = os.path.realpath(TEMP_DIR)

    if not filepath.startswith(temp_dir_real + os.sep):
        raise HTTPException(status_code=400, detail="非法文件名")

    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="文件不存在")

    return FileResponse(
        filepath,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=safe_name,
        background=BackgroundTask(os.remove, filepath),
    )
