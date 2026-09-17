from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "/Users/xionghaoyu/Documents/Deeprevision/DeepRevision项目面试讲稿与高频追问.docx"
BODY_FONT = "STSong"


def set_run_font(run, name=BODY_FONT):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in fonts.attrib:
            del fonts.attrib[key]
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_style_font(style, name=BODY_FONT):
    style.font.name = name
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in fonts.attrib:
            del fonts.attrib[key]
    lang = style._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        style._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_code_line(doc, text):
    p = doc.add_paragraph(style="Code Line")
    p.add_run(text)
    return p


def add_answer(doc, paragraphs, bullets=None):
    for text in paragraphs:
        doc.add_paragraph(text)
    for text in bullets or []:
        add_bullet(doc, text)


_bookmark_id = 1


def add_bookmark(paragraph, name):
    global _bookmark_id
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), str(_bookmark_id))
    bookmark_start.set(qn("w:name"), name)
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), str(_bookmark_id))
    paragraph._p.insert(0, bookmark_start)
    paragraph._p.append(bookmark_end)
    _bookmark_id += 1


def add_internal_link(paragraph, text, anchor):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), BODY_FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([fonts, color, underline])
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_toc_entry(doc, text, anchor, page_number, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28 * (level - 1))
    paragraph.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    run = paragraph.add_run("• " if level == 1 else "  ")
    run.font.bold = level == 1
    add_internal_link(paragraph, text, anchor)
    paragraph.add_run(f"\t{page_number}")
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal)
normal.font.size = Pt(10.8)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.32
normal.paragraph_format.space_after = Pt(6)

for style_name, size, before, after in (
    ("Title", 25, 0, 16),
    ("Heading 1", 17, 18, 8),
    ("Heading 2", 13, 14, 5),
    ("Heading 3", 11.5, 10, 4),
):
    style = styles[style_name]
    set_style_font(style)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title_style = styles["Title"]
title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
title_border = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_border.append(title_bottom)
title_ppr.append(title_border)

if "Code Line" not in styles:
    code_style = styles.add_style("Code Line", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Line"]
set_style_font(code_style)
code_style.font.size = Pt(9.2)
code_style.font.color.rgb = RGBColor(40, 40, 40)
code_style.paragraph_format.left_indent = Inches(0.28)
code_style.paragraph_format.space_before = Pt(1)
code_style.paragraph_format.space_after = Pt(2)
code_style.paragraph_format.line_spacing = 1.12

for list_style in ("List Bullet", "List Bullet 2", "List Number"):
    set_style_font(styles[list_style])
    styles[list_style].font.size = Pt(10.6)
    styles[list_style].paragraph_format.space_after = Pt(3)

header = section.header.paragraphs[0]
header.text = "DeepRevision 项目面试准备"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(header.runs[0])
header.runs[0].font.size = Pt(8.5)
header.runs[0].font.color.rgb = RGBColor(90, 90, 90)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("第 ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)
tail = footer.add_run(" 页")
for r in footer.runs:
    set_run_font(r)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph("DeepRevision 项目面试讲稿与高频追问", style="Title")
subtitle = doc.add_paragraph("十分钟系统陈述与技术追问回答")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(13)
subtitle.runs[0].font.color.rgb = RGBColor(70, 70, 70)

doc.add_paragraph("")
intro = doc.add_paragraph()
intro.add_run("文档用途  ").bold = True
intro.add_run(
    "这份讲稿用于 DeepRevision 项目的面试陈述。第一部分覆盖前端交互、接口代理、Agent 编排、模型接入、RAG、出题组卷、学习闭环和效果验证；第二部分整理十个高频追问及回答；最后补充一次真实出题请求的完整执行链路和 SSE 断开处理。"
)

usage = doc.add_paragraph()
usage.add_run("使用建议  ").bold = True
usage.add_run(
    "先熟悉十分钟陈述的因果顺序，再记住每个追问的首句结论。面试时根据时间删减细节，但保留真实指标、工程边界和生产化改进。"
)

table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.columns[0].width = Inches(1.35)
table.columns[1].width = Inches(2.0)
table.columns[2].width = Inches(3.45)
headers = ["场景", "建议内容", "表达重点"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, "243447")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for r in cell.paragraphs[0].runs:
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
set_repeat_table_header(table.rows[0])
rows = [
    ("三分钟介绍", "讲定位和主链路", "问题、架构、一个核心取舍、一个结果"),
    ("十分钟介绍", "完整使用第一部分", "按一次真实请求的生命周期展开"),
    ("技术追问", "使用第二部分", "先给结论，再解释原因、实现和边界"),
]
for idx, values in enumerate(rows):
    cells = table.add_row().cells
    for i, text in enumerate(values):
        cells[i].text = text
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx % 2 == 1:
            set_cell_shading(cells[i], "F3F6F9")
        for p in cells[i].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(9.3)

doc.add_page_break()
toc_heading = doc.add_paragraph("目录", style="Heading 1")
add_bookmark(toc_heading, "toc")
toc_hint = doc.add_paragraph("点击目录条目可直接跳转；也可以在 Word 的导航窗格中按标题快速定位。")
toc_hint.runs[0].font.color.rgb = RGBColor(90, 90, 90)
for toc_text, toc_anchor, toc_page, toc_level in [
    ("十分钟完整陈述", "speech", 3, 1),
    ("陈述中的关键数字", "metrics", 4, 2),
    ("高频追问与回答", "qa", 4, 1),
    ("问题一 为什么不使用一个大 Agent 调所有工具", "q1", 4, 2),
    ("问题二 为什么还需要 Tool Orchestrator", "q2", 4, 2),
    ("问题三 为什么同时使用 BM25 和向量检索", "q3", 5, 2),
    ("问题四 为什么 RRF 后还需要 Rerank", "q4", 5, 2),
    ("问题五 为什么 HyDE 不默认开启", "q5", 5, 2),
    ("问题六 父子块相比普通分块有什么优势", "q6", 5, 2),
    ("问题七 如何防止模型生成格式错误的题目", "q7", 6, 2),
    ("问题八 mastery 为什么这样计算", "q8", 6, 2),
    ("问题九 当前最大的技术债是什么", "q9", 6, 2),
    ("问题十 如果继续优化会优先做什么", "q10", 6, 2),
    ("面试表达边界", "boundaries", 7, 1),
    ("完整请求执行链路", "request_flow", 8, 1),
    ("SSE 断开与取消链路", "sse_cancel", 9, 1),
]:
    add_toc_entry(doc, toc_text, toc_anchor, toc_page, toc_level)

doc.add_page_break()
speech_heading = doc.add_paragraph("十分钟完整陈述", style="Heading 1")
add_bookmark(speech_heading, "speech")

speech = [
    "我做的是一个面向高校课程复习的 DeepRevision Agent。它解决的不是单纯聊天，而是把课件入库、基于课件问答、出题、作答、掌握度更新、薄弱点复习和试卷导出串成闭环。",
    "前端使用 Next.js 16 和 React 19。用户上传课件后，前端通过 REST 提交文件，再轮询文件状态；聊天使用 POST SSE，前端用 fetch 的 ReadableStream 自己解析 event、data、多行帧和半包，并支持 AbortController 取消。消息协议区分 chat、quiz_set 和 exam_paper，所以普通回答走 Markdown 增量渲染，练习和试卷走结构化组件。前端还能展示引用证据、Agent 路由、工具轨迹和安全确认。",
    "Next.js 同时承担一层 BFF。普通接口通过 catch-all Route Handler 转发给 FastAPI，SSE 使用专门代理，禁用缓存和代理缓冲，并把客户端取消传递给后端。",
    "后端使用 FastAPI。课件上传后先验证 session、后缀和文件路径，再通过 MD5 去重，使用 BackgroundTasks 做解析和向量化。PDF、Word、PPT 和图片分别由对应解析器处理。知识库采用父子索引：子块较小，用于 BM25 和向量检索；父块保存完整章节或连续页面，最终通过 parent_id 回溯。这样解决小块检索准确但上下文不足、大块上下文完整但不容易命中的冲突。",
    "Agent 编排使用 LangGraph。整体采用 Supervisor 集中路由，而不是让多个 Agent 自由对话。高确定性意图先走规则，例如问候、答题跟进、删除确认和生成整卷；开放请求才调用 LLM 输出 route、params 和 reason。Supervisor 后面有 Tool Orchestrator，可以先调用课件搜索、历史查询或管理工具，再把 observation 交给 RAG、Quiz、Exam、Planner、History、Learning Loop、Ops 或 Chitchat Agent。",
    "RAG 链路先用轻量模型做查询改写，同时通过锚点保护避免把课程实体改丢。召回阶段并行执行 BM25 和向量检索，再用加权 RRF 融合，因为两路原始分数不能直接比较。默认先走 fast 候选预算，命中文档少于 3 或来源少于 2 时升级 full；full 后仍低召回才触发 HyDE。RRF 后使用 qwen3.7-text-rerank 精排，设置 8 秒超时，失败就回退到 RRF。检索和重排都在子块上完成，最后回溯父块并限制单来源数量。模型回答后再检查 evidence 是否真实存在于检索上下文，前端展示来源和页码。",
    "模型接入通过统一工厂支持 MiniMax、Token Plan、Qwen 和 DeepSeek，区分主生成模型和轻量模型；调用层处理超时、重试、结构化 JSON 恢复和主备切换。Embedding 使用 Qwen，Rerank 单独接 DashScope，视觉模型是可选能力。",
    "Quiz 和 Exam 不直接信任模型输出，而是执行 Generate、Critique、Revise。题量、题型、编号、选项、答案、解析、重复题和总分先用确定性规则检查，快速模式通过后可以跳过 LLM Critic；完整模式再做语义评审。整卷按选择题、填空判断、简答题分阶段生成，某一阶段失败时只重跑失败阶段，不重新生成整卷。",
    "学生提交答案后，后端把练习记录写入 SQLite，并更新 knowledge_mastery。掌握度由正确率、近期连错和距上次练习时间共同计算，优先复习掌握度低、近期连续错误、长时间未出现的知识点。错题还会进入向量题库，用于相似题复练。Learning Loop Agent 保存 day1_quiz_issued、day1_answered 等状态，确保后续复盘建立在真实练习结果上。",
    "效果验证分为工程、路由、检索、生成和学习效果五层。当前项目测试为 63 passed，前端 production build 通过。RAG 有 15 条人工标注的真实在线基线，Recall@5 为 0.9667，Hit@5 和 MRR 都是 1，nDCG@5 为 0.9585；同时记录了一个复合问题漏召回第二父块的弱案例。生成侧通过质量门和运行指标监控题量、总分、重复、修订和降级，但还没有大规模教育实验，所以不能声称已经证明提升学生成绩。",
    "当前方案适合单机 Demo。生产化会把 SQLite 换成 PostgreSQL、本地课件换成对象存储、BackgroundTasks 换成可恢复任务队列、Chroma 换成独立向量服务，并补充强制认证、租户隔离、限流、Prometheus 和 OpenTelemetry。",
]

for idx, text in enumerate(speech, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.3)
    p.add_run(text)

metrics_heading = doc.add_paragraph("陈述中的关键数字", style="Heading 2")
add_bookmark(metrics_heading, "metrics")
for line in [
    "父子检索：BM25 权重 0.4，向量权重 0.6，RRF 参数 30。",
    "升级策略：fast 命中文档少于 3 或来源少于 2 时升级 full。",
    "外部重排：qwen3.7-text-rerank，超时约 8 秒后回退 RRF。",
    "当前验证：63 个测试通过，前端生产构建通过。",
    "RAG 基线：15 条人工标注查询，Recall@5 为 0.9667，MRR 为 1。",
]:
    add_bullet(doc, line)

qa_heading = doc.add_paragraph("高频追问与回答", style="Heading 1")
add_bookmark(qa_heading, "qa")
doc.add_paragraph(
    "回答追问时，先用一两句话给出结论，再根据面试官兴趣展开原因、实现和局限。以下内容可直接作为回答底稿。"
)

questions = [
    (
        "问题一 为什么不使用一个大 Agent 调所有工具",
        [
            "课程问答、出题、整卷和数据管理的输出约束不同。把所有能力放进一个大 Agent，会让 prompt 越来越复杂，路由结果、调用成本和错误原因都更难定位。",
            "这个项目让 Supervisor 显式选择专项 Agent。每个 Agent 只处理自己的输入输出契约，同时把问候、作答跟进和危险操作确认交给确定性规则，减少 LLM 路由抖动。",
        ],
        ["专项 Agent 更容易单独测试和评测。", "高风险操作可以在统一入口设置安全闸。", "不同任务可以使用不同超时、模型和质量门。"],
    ),
    (
        "问题二 为什么还需要 Tool Orchestrator",
        [
            "Supervisor 解决把请求交给谁，Tool Orchestrator 解决交给专项 Agent 前需要执行哪些工具。两者负责不同层次的决策。",
            "例如 RAG 回答前可能先检索课件，Planner 生成计划前可能先读取 mastery。Orchestrator 会保存工具 observation，后续 Agent 可以复用结果，避免重复查询；危险操作也在这一层被拦截。",
        ],
        ["工具结果可以进入执行轨迹。", "Orchestrator 可以根据 observation 调整最终目标 Agent。", "工具上下文可跨轮保存，但需要限制大小和保留时间。"],
    ),
    (
        "问题三 为什么同时使用 BM25 和向量检索",
        [
            "BM25 对课程术语、缩写和精确关键词更稳定，向量检索对同义表达和自然语言改写更稳定。课程课件同时具有这两类特征，单一路径容易漏召回。",
            "项目用加权 RRF 融合两路排名。RRF 依赖名次，不要求 BM25 分数和向量相似度位于同一数值尺度。",
        ],
        ["当前权重为 BM25 0.4、向量 0.6。", "先优化召回，再用 Rerank 优化前几名排序。", "权重需要在固定评测集上调，而不是凭感觉设置。"],
    ),
    (
        "问题四 为什么 RRF 后还需要 Rerank",
        [
            "RRF 的主要任务是融合召回结果。Cross Encoder Rerank 同时阅读 query 和候选文档，更适合判断候选之间的细粒度相关性，因此能够改善 MRR、nDCG 和前几名精度。",
            "Rerank 不能替代召回，因为没有进入候选集的文档无法被重排。为了控制成本，项目只对 RRF 缩小后的候选集调用 qwen3.7-text-rerank，并在超时或接口失败时保留 RRF 顺序。",
        ],
        ["召回解决找得到，重排解决排得准。", "重排超时约 8 秒，采用 fail open。", "是否值得重排应结合延迟和排序指标评估。"],
    ),
    (
        "问题五 为什么 HyDE 不默认开启",
        [
            "HyDE 需要额外调用一次模型，并用生成的假设教材片段发起向量检索。它能帮助短问题或关键词不足的问题，但也会增加延迟，并可能让假设内容影响检索方向。",
            "因此项目只在 full 检索后仍然低召回时触发 HyDE，把额外成本集中在困难查询上。",
        ],
        ["HyDE 只用于检索，不直接作为最终答案证据。", "触发次数和补召回数量应进入 Trace。", "若真实评测没有收益，应关闭或调整触发阈值。"],
    ),
    (
        "问题六 父子块相比普通分块有什么优势",
        [
            "父子块把检索粒度和生成粒度解耦。子块较小，便于准确匹配；父块较大，保留章节、连续页面和完整论述。",
            "系统先在子块上完成 BM25、向量检索和 Rerank，再通过 parent_id 回溯父块。这样不必在小块准确和大块完整之间二选一。",
        ],
        ["父块需要独立存储并维护与子块的一致性。", "删除或重建课件时必须同步清理父块和子块。", "块大小需要通过真实评测校准。"],
    ),
    (
        "问题七 如何防止模型生成格式错误的题目",
        [
            "项目同时使用结构化 schema 和生成后质量门。schema 约束字段形态，质量门继续验证题量、题型、编号、选项、答案、解析、重复和总分。",
            "可以由规则确定的条件优先用本地检查。只有歧义、答案合理性和课件一致性等语义问题才调用 LLM Critic。修订轮数和总耗时都有上限，失败时记录 delivery_mode 和 degrade_reason。",
        ],
        ["快速模式允许本地检查通过后跳过 LLM Critic。", "完整模式强制增加语义评审。", "底线质量仍不满足时应拒绝交付。"],
    ),
    (
        "问题八 mastery 为什么这样计算",
        [
            "当前目标是快速形成可解释的学习闭环，因此 mastery 使用正确率作为基础，再扣除近期连错和遗忘惩罚。这个分数主要用于安排复习优先级。",
            "它不是经过教育测量验证的认知诊断模型。获得足够的学生交互数据后，可以比较 IRT、BKT、DKT 或基于知识图谱的学习状态模型。",
        ],
        ["连续错误每次扣 0.12，最高扣 0.36。", "距上次练习每天扣 0.015，最高扣 0.20。", "review urgency 等于 1 减去动态 mastery。"],
    ),
    (
        "问题九 当前最大的技术债是什么",
        [
            "当前架构的主要技术债来自单机 Demo 假设。FastAPI BackgroundTasks 在进程重启后不能恢复，本地 SQLite 和 Chroma 不适合真正的多实例部署，运行指标也保存在进程内存中。",
            "此外，练习提交的幂等键还需要前端完整接入，生成质量缺少规模化人工评测，学习效果也没有 A B 实验。",
        ],
        ["任务恢复和数据存储应优先改造。", "生产指标需要外部持久化和跨实例聚合。", "不能把检索离线指标等同于学习效果。"],
    ),
    (
        "问题十 如果继续优化会优先做什么",
        [
            "我会先处理可靠性和可评测性，再增加更复杂的 Agent 能力。第一步是把入库迁移到可恢复任务队列，并把 SQLite、本地课件和本地向量库迁移为 PostgreSQL、对象存储和独立向量服务。",
            "第二步是冻结真实检索和出题评测集，针对复合问题引入 query decomposition 和 multi query retrieval。随后补齐 OpenTelemetry、Prometheus、认证、租户隔离和幂等性，最后再用真实学生数据校准 mastery。",
        ],
        [
            "可恢复任务队列和持久化任务状态。",
            "真实评测集和版本化实验报告。",
            "复合问题拆解与多查询召回。",
            "PostgreSQL、对象存储和独立向量服务。",
            "OpenTelemetry、Prometheus、限流和审计。",
            "基于真实数据校准学习状态模型。",
        ],
    ),
]

for question_index, (heading, paragraphs, bullets) in enumerate(questions, start=1):
    question_heading = doc.add_paragraph(heading, style="Heading 2")
    add_bookmark(question_heading, f"q{question_index}")
    add_answer(doc, paragraphs, bullets)

boundary_heading = doc.add_paragraph("面试表达边界", style="Heading 1")
add_bookmark(boundary_heading, "boundaries")
for text in [
    "最新真实检索基线是 15 条人工标注查询。仓库还保留过旧版 8 条和模拟 10 条实验，面试时不要混用。",
    "可以说学习数据闭环已经实现并通过工程测试，不能说已经通过大规模教育实验显著提高成绩。",
    "可以说系统支持可选认证、会话 ACL 和幂等接口，不能说默认本地演示已经具备完整多租户生产安全体系。",
    "可以说模型调用具有超时、重试和主备切换，不能承诺仓库没有提供的并发、可用性或成本 SLA。",
]:
    add_bullet(doc, text)

doc.add_page_break()
flow_heading = doc.add_paragraph("完整请求执行链路", style="Heading 1")
add_bookmark(flow_heading, "request_flow")
doc.add_paragraph(
    "下面以用户输入“根据最近错题针对薄弱知识点出 5 道题”为例，按一次真实请求从前端到学习闭环的生命周期展开。"
)

doc.add_paragraph("前端交互与请求转发", style="Heading 2")
doc.add_paragraph(
    "React 前端先做空输入和重复提交保护，并乐观插入用户消息与空的助手消息。请求通过 Next.js Route Handler 转发到 FastAPI，同时保持 SSE 流和取消信号。"
)

doc.add_paragraph("身份校验与路由决策", style="Heading 2")
doc.add_paragraph(
    "后端先完成 Bearer 身份、Scope 和 Session ACL 校验，再按 session_id 从 SQLite 读取近期对话、长期记忆和练习画像，构造 SupervisorState。"
)
doc.add_paragraph(
    "Supervisor 采用规则和 LLM 混合路由。这句话没有明确的复习计划或闭环表达，因此主路由是 Quiz，而不是 Learning Loop。Router 同时提取 topic、题型和数量；输出要经过路由白名单、参数范围和语义漂移校验。模型失败时，系统使用“出题”关键词兜底。随后，Tool Orchestrator 对复杂请求执行有限步 ReAct 规划。工具按只读、受控写、危险写和委派分级，危险操作必须二次确认。"
)

doc.add_paragraph("薄弱知识点识别", style="Heading 2")
doc.add_paragraph(
    "QuizAgent 不从聊天文本猜测薄弱点，而是直接读取 practice_records 和 knowledge_mastery。掌握度综合正确率、连续错误和时间衰减，得到复习紧迫度，再把优先薄弱点加入出题主题。同时，系统读取最近 15 轮题目签名，避免重复出题。"
)

doc.add_paragraph("课件检索与父子块回溯", style="Heading 2")
doc.add_paragraph(
    "出题前通过 RAG 获取课件依据。系统将课件切成小的检索子块和大的生成父块，子块进入 BM25 与 Qwen Embedding 索引。检索时先执行 Query Rewrite，再并行执行 BM25 和向量召回，通过 0.4 和 0.6 加权的 RRF 融合。默认先走 8 条候选的 fast 档；召回数量或来源不足时升级到 16 条 full 档；full 档仍然低召回时再触发 HyDE。之后使用 qwen3.7-text-rerank 精排候选，并根据 parent_id 回溯完整父块。"
)

doc.add_paragraph("题目生成质量门禁与学习闭环", style="Heading 2")
doc.add_paragraph(
    "拿到课件上下文后，QuizAgent 执行 Generate、Critic、Revise 的 Reflexion 工作流。生成结果经过 Pydantic、数量、选项、答案一致性、重复度和证据相关性等质量门禁，达不到底线就降级或拒绝交付。"
)
doc.add_paragraph(
    "最终结果以 quiz_set 和 interactive_cards 结构通过 SSE 返回前端。前端解析 start、heartbeat、complete 和 done 事件，再把结构化 payload 渲染成题卡。用户提交答案后，系统通过幂等接口写入练习记录并更新掌握度。下一轮出题继续使用更新后的画像，从而形成课件、练习、画像和再练习的闭环。"
)

sse_heading = doc.add_paragraph("SSE 断开与取消链路", style="Heading 1")
add_bookmark(sse_heading, "sse_cancel")
doc.add_paragraph("SSE 断开后，取消信号按下面的顺序向后传播。")
for step in [
    "前端通过 AbortController 发出取消信号。",
    "Next.js 感知下游连接断开，并取消上游 FastAPI 请求。",
    "FastAPI 通过 request disconnect 或 cancellation 感知中断。",
    "后端设置任务取消状态，并把取消信号传播到 Agent Runtime。",
    "Agent Runtime 中断 LLM 流，停止后续 Tool 调用，并阻止尚未开始的写库操作。",
    "finally 代码块统一释放流、网络连接和任务资源。",
]:
    add_number(doc, step)

doc.core_properties.title = "DeepRevision 项目面试讲稿与高频追问"
doc.core_properties.subject = "十分钟系统陈述与技术追问回答"
doc.core_properties.author = "DeepRevision 项目组"
doc.core_properties.keywords = "DeepRevision Agent RAG 面试"

# LibreOffice and Word can choose different East Asian fallbacks when a run has
# no direct font. Pin every authored run to a font that contains Simplified
# Chinese glyphs so the render verification matches the delivered document.
for paragraph in doc.paragraphs:
    for authored_run in paragraph.runs:
        set_run_font(authored_run)
for authored_table in doc.tables:
    for row in authored_table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for authored_run in paragraph.runs:
                    set_run_font(authored_run)

doc.save(OUTPUT)
print(OUTPUT)
